import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(11, 19, 43)      # #0B132B Deep Space Navy
    CARD_BG = RGBColor(21, 32, 64)      # #152040 Card fill
    CARD_BORDER = RGBColor(45, 62, 105) # Border
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(160, 174, 192)
    CYAN_ACCENT = RGBColor(0, 229, 255)  # Tech cyan
    GREEN_ACCENT = RGBColor(16, 185, 129)# Emerald green
    GOLD_ACCENT = RGBColor(245, 158, 11) # Warning/Highlight gold
    ORANGE_ACCENT = RGBColor(249, 115, 22)# Warm accent
    PURPLE_ACCENT = RGBColor(139, 92, 246)

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()

    def add_header(slide, title_text, category_badge="SMART INDIA HACKATHON 2026"):
        # Header Top Bar
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Category/Hackathon badge
        p_sub = tf.paragraphs[0]
        p_sub.text = category_badge.upper()
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = CYAN_ACCENT
        p_sub.space_after = Pt(2)
        
        # Main Slide Title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Team Logo Badge on top right
        team_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.8), Inches(0.4), Inches(1.7), Inches(0.6))
        team_badge.fill.solid()
        team_badge.fill.fore_color.rgb = CARD_BG
        team_badge.line.color.rgb = CYAN_ACCENT
        team_badge.line.width = Pt(1.5)
        tf_tb = team_badge.text_frame
        p_tb = tf_tb.paragraphs[0]
        p_tb.text = "Team GeoNex"
        p_tb.font.size = Pt(12)
        p_tb.font.bold = True
        p_tb.font.color.rgb = CYAN_ACCENT
        p_tb.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 1: TITLE & OVERVIEW
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1)
    
    # Hero Title Container
    hero_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(2.2))
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = CARD_BG
    hero_card.line.color.rgb = CYAN_ACCENT
    hero_card.line.width = Pt(2)
    
    tf1 = hero_card.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.4)
    tf1.margin_top = Inches(0.3)

    p1 = tf1.paragraphs[0]
    p1.text = "SMART INDIA HACKATHON 2026"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = GOLD_ACCENT

    p2 = tf1.add_paragraph()
    p2.text = "AI-Powered Drone-Based Geospatial Mapping & Change Detection System"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Automated Urban Parcel Mapping & Cadastral Feature Extraction from High-Resolution Aerial Imagery"
    p3.font.size = Pt(13)
    p3.font.color.rgb = CYAN_ACCENT
    p3.space_before = Pt(6)

    # 4 Metadata Cards Grid below hero
    meta_items = [
        ("PROBLEM STATEMENT ID", "26012", CYAN_ACCENT),
        ("THEME", "Smart Automation", GREEN_ACCENT),
        ("PS CATEGORY", "Software", GOLD_ACCENT),
        ("TEAM NAME", "GeoNex", ORANGE_ACCENT)
    ]

    card_width = Inches(2.7)
    card_gap = Inches(0.31)
    top_pos = Inches(3.1)

    for i, (label, val, color) in enumerate(meta_items):
        left_pos = Inches(0.8) + i * (card_width + card_gap)
        mcard = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_width, Inches(1.3))
        mcard.fill.solid()
        mcard.fill.fore_color.rgb = CARD_BG
        mcard.line.color.rgb = CARD_BORDER
        mcard.line.width = Pt(1)

        mtf = mcard.text_frame
        mtf.margin_left = Inches(0.2)
        mtf.margin_top = Inches(0.2)
        
        mp1 = mtf.paragraphs[0]
        mp1.text = label
        mp1.font.size = Pt(10)
        mp1.font.bold = True
        mp1.font.color.rgb = TEXT_MUTED

        mp2 = mtf.add_paragraph()
        mp2.text = val
        mp2.font.size = Pt(18)
        mp2.font.bold = True
        mp2.font.color.rgb = color
        mp2.space_before = Pt(4)

    # Full PS Description Banner
    ps_banner = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.7), Inches(11.733), Inches(2.2))
    ps_banner.fill.solid()
    ps_banner.fill.fore_color.rgb = CARD_BG
    ps_banner.line.color.rgb = CARD_BORDER
    
    btf = ps_banner.text_frame
    btf.margin_left = Inches(0.4)
    btf.margin_top = Inches(0.3)
    btf.word_wrap = True

    bp1 = btf.paragraphs[0]
    bp1.text = "PROBLEM STATEMENT DETAILS"
    bp1.font.size = Pt(12)
    bp1.font.bold = True
    bp1.font.color.rgb = CYAN_ACCENT

    bp2 = btf.add_paragraph()
    bp2.text = "Title: AI-Based Automated Urban Parcel Mapping and Cadastral Feature Extraction System using Drone Imagery."
    bp2.font.size = Pt(14)
    bp2.font.bold = True
    bp2.font.color.rgb = TEXT_WHITE
    bp2.space_before = Pt(6)

    bp3 = btf.add_paragraph()
    bp3.text = "Core Challenge: Conventional urban cadastral mapping relies on labor-intensive manual survey methods or full re-processing of aerial imagery for every update cycle. GeoNex solves this by combining UNet++ multi-class semantic segmentation with automated vectorization and incremental AI change detection."
    bp3.font.size = Pt(12)
    bp3.font.color.rgb = TEXT_MUTED
    bp3.space_before = Pt(8)


    # ==========================================
    # SLIDE 2: PROBLEM & SOLUTION ARCHITECTURE
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_bg(slide2)
    add_header(slide2, "Problem & Proposed Solution Overview")

    # Left Column: Process Flow (Horizontal Cards stacked)
    left_title = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(0.4))
    ltf = left_title.text_frame
    lp = ltf.paragraphs[0]
    lp.text = "END-TO-END PIPELINE WORKFLOW"
    lp.font.size = Pt(13)
    lp.font.bold = True
    lp.font.color.rgb = CYAN_ACCENT

    pipeline_steps = [
        ("1. Drone Imagery", "Raw GeoTIFF / ECW orthomosaics", CYAN_ACCENT),
        ("2. Pre-Processing", "Tiling, normalization & georeferencing", TEXT_WHITE),
        ("3. AI Segmentation", "UNet++ (Buildings, Roads, Water, Rooftops)", GREEN_ACCENT),
        ("4. GIS Vectorization", "Raster masks → GeoJSON / Shapefiles", GOLD_ACCENT),
        ("5. Master Spatial DB", "Centralized PostGIS spatial repository", ORANGE_ACCENT),
        ("6. Incremental Update", "Detect changes & update only modified parcels", CYAN_ACCENT)
    ]

    p_top = Inches(1.95)
    for i, (title, desc, accent) in enumerate(pipeline_steps):
        scard = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), p_top + i * Inches(0.85), Inches(5.6), Inches(0.75))
        scard.fill.solid()
        scard.fill.fore_color.rgb = CARD_BG
        scard.line.color.rgb = accent
        scard.line.width = Pt(1)

        stf = scard.text_frame
        stf.margin_left = Inches(0.25)
        stf.margin_top = Inches(0.12)
        
        stp1 = stf.paragraphs[0]
        stp1.text = title
        stp1.font.size = Pt(12)
        stp1.font.bold = True
        stp1.font.color.rgb = accent

        stp2 = stf.add_paragraph()
        stp2.text = desc
        stp2.font.size = Pt(10)
        stp2.font.color.rgb = TEXT_MUTED

    # Right Column: How it Addresses Problem & Innovations
    rcard1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.95), Inches(5.833), Inches(2.5))
    rcard1.fill.solid()
    rcard1.fill.fore_color.rgb = CARD_BG
    rcard1.line.color.rgb = CARD_BORDER

    rtf1 = rcard1.text_frame
    rtf1.margin_left = rtf1.margin_top = Inches(0.3)
    rtf1.word_wrap = True

    rp1 = rtf1.paragraphs[0]
    rp1.text = "HOW IT ADDRESSES THE PROBLEM"
    rp1.font.size = Pt(13)
    rp1.font.bold = True
    rp1.font.color.rgb = GREEN_ACCENT

    bullets1 = [
        "Automates extraction of geographic features from drone imagery.",
        "Reduces manual GIS digitizing effort by over 80%.",
        "Maintains a reusable master geospatial database.",
        "New surveys update only changed regions instead of reprocessing everything.",
        "Delivers GIS map-ready outputs instantly for urban planning."
    ]
    for b in bullets1:
        bp = rtf1.add_paragraph()
        bp.text = f"•  {b}"
        bp.font.size = Pt(11)
        bp.font.color.rgb = TEXT_WHITE
        bp.space_before = Pt(4)

    rcard2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(4.6), Inches(5.833), Inches(2.45))
    rcard2.fill.solid()
    rcard2.fill.fore_color.rgb = CARD_BG
    rcard2.line.color.rgb = CARD_BORDER

    rtf2 = rcard2.text_frame
    rtf2.margin_left = rtf2.margin_top = Inches(0.3)
    rtf2.word_wrap = True

    rp2 = rtf2.paragraphs[0]
    rp2.text = "INNOVATION & KEY UNIQUENESS"
    rp2.font.size = Pt(13)
    rp2.font.bold = True
    rp2.font.color.rgb = GOLD_ACCENT

    bullets2 = [
        "Incremental Geospatial Updating: Only recompute modified land parcels.",
        "AI-Based Change Detection: Pixel & vector difference engine.",
        "Multi-Class Feature Extraction: Buildings + Roads + Water + Rooftop types.",
        "Confidence Validation: Human-in-the-loop validation for low-confidence AI masks.",
        "Interactive Web Dashboard: Real-time map overlay and spatial analytics."
    ]
    for b in bullets2:
        bp = rtf2.add_paragraph()
        bp.text = f"•  {b}"
        bp.font.size = Pt(11)
        bp.font.color.rgb = TEXT_WHITE
        bp.space_before = Pt(4)


    # ==========================================
    # SLIDE 3: TECHNICAL APPROACH & ARCHITECTURE
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_bg(slide3)
    add_header(slide3, "Technical Approach & System Architecture")

    # 6 Stage Pipeline Boxes Across Top
    pipe_stages = [
        ("1. Input Data", "Drone Orthomosaic\n(ECW / GeoTIFF)", CYAN_ACCENT),
        ("2. Pre-Processing", "Tiling, Resizing,\nGeo-Referencing", TEXT_WHITE),
        ("3. AI Segmentation", "UNet++ Deep Model\n(Dice + CE Loss)", GREEN_ACCENT),
        ("4. Feature Extract", "Noise Removal &\nPolygonization", GOLD_ACCENT),
        ("5. GIS Mapping", "GeoJSON Output &\nWeb Visualizer", ORANGE_ACCENT),
        ("6. Change Engine", "Difference Analysis &\nPostGIS Update", PURPLE_ACCENT)
    ]

    box_w = Inches(1.82)
    box_gap = Inches(0.16)
    for i, (stitle, sdesc, scolor) in enumerate(pipe_stages):
        bleft = Inches(0.8) + i * (box_w + box_gap)
        pbox = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bleft, Inches(1.5), box_w, Inches(2.2))
        pbox.fill.solid()
        pbox.fill.fore_color.rgb = CARD_BG
        pbox.line.color.rgb = scolor
        pbox.line.width = Pt(1.5)

        ptf = pbox.text_frame
        ptf.margin_left = ptf.margin_right = ptf.margin_top = Inches(0.15)
        ptf.word_wrap = True

        pp1 = ptf.paragraphs[0]
        pp1.text = stitle
        pp1.font.size = Pt(11)
        pp1.font.bold = True
        pp1.font.color.rgb = scolor

        pp2 = ptf.add_paragraph()
        pp2.text = sdesc
        pp2.font.size = Pt(10)
        pp2.font.color.rgb = TEXT_MUTED
        pp2.space_before = Pt(8)

    # Lower Section: Architecture Tech Stack Grid (4 Stack Columns)
    arch_title = slide3.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.733), Inches(0.35))
    atf = arch_title.text_frame
    ap = atf.paragraphs[0]
    ap.text = "SYSTEM ARCHITECTURE & TECH STACK"
    ap.font.size = Pt(13)
    ap.font.bold = True
    ap.font.color.rgb = CYAN_ACCENT

    tech_stacks = [
        ("Frontend (Web App)", ["Interactive Map (Leaflet / Mapbox)", "Layer Toggle (Buildings/Roads)", "Result Analytics Dashboard"], CYAN_ACCENT),
        ("Backend (FastAPI)", ["RESTful API Endpoints", "Asynchronous Processing Pipeline", "GeoJSON & Vector Exporter"], GREEN_ACCENT),
        ("Database (PostGIS)", ["PostgreSQL + PostGIS Extension", "Spatial Parcel Indexing", "Survey Versioning System"], GOLD_ACCENT),
        ("Cloud & Compute", ["AWS S3 Image Store", "GPU Inference (PyTorch)", "Docker Microservices"], PURPLE_ACCENT)
    ]

    col_w = Inches(2.78)
    col_gap = Inches(0.2)
    for i, (stitle, sitems, scolor) in enumerate(tech_stacks):
        cleft = Inches(0.8) + i * (col_w + col_gap)
        abox = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cleft, Inches(4.35), col_w, Inches(2.6))
        abox.fill.solid()
        abox.fill.fore_color.rgb = CARD_BG
        abox.line.color.rgb = CARD_BORDER

        abtf = abox.text_frame
        abtf.margin_left = abtf.margin_top = Inches(0.2)
        abtf.word_wrap = True

        ap1 = abtf.paragraphs[0]
        ap1.text = stitle
        ap1.font.size = Pt(12)
        ap1.font.bold = True
        ap1.font.color.rgb = scolor

        for item in sitems:
            ap2 = abtf.add_paragraph()
            ap2.text = f"• {item}"
            ap2.font.size = Pt(10)
            ap2.font.color.rgb = TEXT_WHITE
            ap2.space_before = Pt(6)


    # ==========================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_bg(slide4)
    add_header(slide4, "Feasibility, Viability & Risk Mitigation")

    # Top Left 3 Pillar Cards
    pillars = [
        ("Technical Feasibility", "Leverages established AI models (UNet++), GIS engines (GDAL/Rasterio), and standard cloud GPUs.", CYAN_ACCENT),
        ("Operational Viability", "Initial survey establishes master parcel DB; subsequent surveys process incremental updates seamlessly.", GREEN_ACCENT),
        ("Scalability", "Tiled image processing + Dockerized microservices enable scaling to state-wide satellite & drone imagery.", GOLD_ACCENT)
    ]

    pw = Inches(3.7)
    pgap = Inches(0.3)
    for i, (title, desc, color) in enumerate(pillars):
        pleft = Inches(0.8) + i * (pw + pgap)
        pcard = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pleft, Inches(1.5), pw, Inches(1.7))
        pcard.fill.solid()
        pcard.fill.fore_color.rgb = CARD_BG
        pcard.line.color.rgb = color
        pcard.line.width = Pt(1.5)

        ptf = pcard.text_frame
        ptf.margin_left = ptf.margin_top = Inches(0.2)
        ptf.word_wrap = True

        pp1 = ptf.paragraphs[0]
        pp1.text = title
        pp1.font.size = Pt(12)
        pp1.font.bold = True
        pp1.font.color.rgb = color

        pp2 = ptf.add_paragraph()
        pp2.text = desc
        pp2.font.size = Pt(10)
        pp2.font.color.rgb = TEXT_WHITE
        pp2.space_before = Pt(6)

    # Bottom Table Container: Technical Challenges & Proposed Solutions
    t_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.4), Inches(11.733), Inches(3.6))
    t_box.fill.solid()
    t_box.fill.fore_color.rgb = CARD_BG
    t_box.line.color.rgb = CARD_BORDER

    ttf = t_box.text_frame
    ttf.margin_left = ttf.margin_top = Inches(0.25)
    ttf.word_wrap = True

    tp1 = ttf.paragraphs[0]
    tp1.text = "TECHNICAL CHALLENGES & PROPOSED SOLUTIONS MATRIX"
    tp1.font.size = Pt(13)
    tp1.font.bold = True
    tp1.font.color.rgb = CYAN_ACCENT

    table_data = [
        ("Limited Training Data", "Data augmentation (flips, rotations, spectral shifts) + transfer learning from aerial datasets."),
        ("Building Segmentation Errors", "Improved UNet++ loss formulation (Dice + Focal Loss) + morphological boundary smoothing."),
        ("Water / Algae Confusion", "Hard-negative mining during training + multi-spectral NDWI band integration where available."),
        ("Image Misalignment & Large Size", "Automated geospatial registration + adaptive overlapping tile processing."),
        ("Tile Boundary Artifacts", "Overlap-aware stitching with fuzzy boundary weight merging."),
        ("Reprocessing Entire Survey Area", "AI-based vector difference engine: recomputes ONLY changed land parcels (saving >80% compute).")
    ]

    for chal, sol in table_data:
        p = ttf.add_paragraph()
        p.space_before = Pt(5)
        # Bold Challenge, normal Solution
        r1 = p.add_run()
        r1.text = f"•  {chal}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = GOLD_ACCENT

        r2 = p.add_run()
        r2.text = sol
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_WHITE


    # ==========================================
    # SLIDE 5: IMPACT AND BENEFITS (VISUAL COMPARISON)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_bg(slide5)
    add_header(slide5, "Impact, Benefits & Visual Process Shift")

    # Left Box: TRADITIONAL vs GEONEX PROPOSED Visual Comparison
    comp_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.3))
    comp_card.fill.solid()
    comp_card.fill.fore_color.rgb = CARD_BG
    comp_card.line.color.rgb = CARD_BORDER

    ctf = comp_card.text_frame
    ctf.margin_left = ctf.margin_top = Inches(0.25)
    ctf.word_wrap = True

    cp1 = ctf.paragraphs[0]
    cp1.text = "WORKFLOW COMPARISON (PARADIGM SHIFT)"
    cp1.font.size = Pt(13)
    cp1.font.bold = True
    cp1.font.color.rgb = CYAN_ACCENT

    # Traditional sub box
    trad_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(2.1), Inches(5.0), Inches(1.5))
    trad_box.fill.solid()
    trad_box.fill.fore_color.rgb = RGBColor(40, 20, 30)
    trad_box.line.color.rgb = RGBColor(239, 68, 68)

    ttf = trad_box.text_frame
    ttf.margin_left = ttf.margin_top = Inches(0.15)
    tp = ttf.paragraphs[0]
    tp.text = "TRADITIONAL METHOD (Inefficient)"
    tp.font.size = Pt(11)
    tp.font.bold = True
    tp.font.color.rgb = RGBColor(248, 113, 113)

    t_flow = ttf.add_paragraph()
    t_flow.text = "Full Area Image  ➔  Full AI Processing  ➔  Repeat Every Survey  ➔  High Compute & Slow Update"
    t_flow.font.size = Pt(10)
    t_flow.font.color.rgb = TEXT_MUTED
    t_flow.space_before = Pt(6)

    # Proposed sub box
    prop_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(3.8), Inches(5.0), Inches(1.8))
    prop_box.fill.solid()
    prop_box.fill.fore_color.rgb = RGBColor(16, 45, 40)
    prop_box.line.color.rgb = GREEN_ACCENT
    prop_box.line.width = Pt(1.5)

    ptf = prop_box.text_frame
    ptf.margin_left = ptf.margin_top = Inches(0.15)
    pp = ptf.paragraphs[0]
    pp.text = "GEONEX PROPOSED METHOD (Smart & Fast)"
    pp.font.size = Pt(11)
    pp.font.bold = True
    pp.font.color.rgb = GREEN_ACCENT

    p_flow = ptf.add_paragraph()
    p_flow.text = "Existing Cadastral Map + New Drone Survey\n   ↓\nAI Change Detection Engine\n   ↓\nRe-process CHANGED PARCELS ONLY  ➔ Instant Database Update"
    p_flow.font.size = Pt(10)
    p_flow.font.color.rgb = TEXT_WHITE
    p_flow.space_before = Pt(6)

    # Right Box: Categorized Key Benefits
    bcard = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.5), Inches(5.833), Inches(4.3))
    bcard.fill.solid()
    bcard.fill.fore_color.rgb = CARD_BG
    bcard.line.color.rgb = CARD_BORDER

    btf = bcard.text_frame
    btf.margin_left = btf.margin_top = Inches(0.25)
    btf.word_wrap = True

    bp1 = btf.paragraphs[0]
    bp1.text = "KEY BENEFIT CATEGORIES"
    bp1.font.size = Pt(13)
    bp1.font.bold = True
    bp1.font.color.rgb = GOLD_ACCENT

    b_categories = [
        ("Planning & Cadastral", "Rapid building & road extraction for urban growth monitoring.", CYAN_ACCENT),
        ("Economic Efficiency", "Eliminates repetitive manual digitizing & reduces compute costs by >80%.", GREEN_ACCENT),
        ("Environmental Monitoring", "Tracks water body encroachments, vegetation changes, and land-use shifts.", GOLD_ACCENT),
        ("Disaster Management", "Rapid post-disaster damage assessment and updated hazard maps.", PURPLE_ACCENT)
    ]
    for cat, desc, accent in b_categories:
        bp = btf.add_paragraph()
        bp.space_before = Pt(8)
        r1 = bp.add_run()
        r1.text = f"•  {cat}: "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = accent

        r2 = bp.add_run()
        r2.text = desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_WHITE

    # Bottom Quote Punchline Banner
    quote_banner = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.733), Inches(1.0))
    quote_banner.fill.solid()
    quote_banner.fill.fore_color.rgb = CARD_BG
    quote_banner.line.color.rgb = GOLD_ACCENT
    quote_banner.line.width = Pt(1.5)

    qtf = quote_banner.text_frame
    qtf.margin_left = Inches(0.3)
    qtf.margin_top = Inches(0.2)

    qp = qtf.paragraphs[0]
    qp.text = "“Build the map once, detect the changes later, and update only what has changed.”"
    qp.font.size = Pt(15)
    qp.font.bold = True
    qp.font.italic = True
    qp.font.color.rgb = GOLD_ACCENT
    qp.alignment = PP_ALIGN.CENTER


    # ==========================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_bg(slide6)
    add_header(slide6, "Research Areas & Key References")

    # 3 Column Layout
    col_width3 = Inches(3.7)
    gap3 = Inches(0.31)

    # Col 1: Primary References & Code
    rbox1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), col_width3, Inches(5.3))
    rbox1.fill.solid()
    rbox1.fill.fore_color.rgb = CARD_BG
    rbox1.line.color.rgb = CYAN_ACCENT

    rtf1 = rbox1.text_frame
    rtf1.margin_left = rtf1.margin_top = Inches(0.25)
    rtf1.word_wrap = True

    rp1 = rtf1.paragraphs[0]
    rp1.text = "PRIMARY REPOSITORY & REFS"
    rp1.font.size = Pt(13)
    rp1.font.bold = True
    rp1.font.color.rgb = CYAN_ACCENT

    col1_items = [
        ("GitHub Repository", "https://github.com/Kabeer2004/ProjectVaayu"),
        ("Drone Feature Extraction", "High-resolution UAV orthomosaic processing pipeline"),
        ("UNet++ Segmentation", "Nested U-Net architecture for multi-class parcel extraction"),
        ("GDAL / GIS Pipeline", "Rasterio & GDAL spatial transformation workflows")
    ]
    for head, body in col1_items:
        p = rtf1.add_paragraph()
        p.space_before = Pt(12)
        r1 = p.add_run()
        r1.text = f"{head}\n"
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_WHITE

        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # Col 2: Core Research Areas
    rbox2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_width3 + gap3, Inches(1.5), col_width3, Inches(5.3))
    rbox2.fill.solid()
    rbox2.fill.fore_color.rgb = CARD_BG
    rbox2.line.color.rgb = GREEN_ACCENT

    rtf2 = rbox2.text_frame
    rtf2.margin_left = rtf2.margin_top = Inches(0.25)
    rtf2.word_wrap = True

    rp2 = rtf2.paragraphs[0]
    rp2.text = "CORE RESEARCH DOMAINS"
    rp2.font.size = Pt(13)
    rp2.font.bold = True
    rp2.font.color.rgb = GREEN_ACCENT

    col2_items = [
        "Semantic Segmentation in High-Res Remote Sensing",
        "Deep Learning for Cadastral Feature Extraction",
        "Automated Vectorization (Raster to GeoJSON)",
        "Bi-temporal Spatial Change Detection",
        "Geospatial Image Registration & Alignment",
        "PostGIS Spatial Indexing & Database Versioning"
    ]
    for item in col2_items:
        p = rtf2.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)

    # Col 3: Technology Stack & Libraries
    rbox3 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + 2*(col_width3 + gap3), Inches(1.5), col_width3, Inches(5.3))
    rbox3.fill.solid()
    rbox3.fill.fore_color.rgb = CARD_BG
    rbox3.line.color.rgb = GOLD_ACCENT

    rtf3 = rbox3.text_frame
    rtf3.margin_left = rtf3.margin_top = Inches(0.25)
    rtf3.word_wrap = True

    rp3 = rtf3.paragraphs[0]
    rp3.text = "TECH STACK & FRAMEWORKS"
    rp3.font.size = Pt(13)
    rp3.font.bold = True
    rp3.font.color.rgb = GOLD_ACCENT

    col3_items = [
        ("AI / ML Models", "PyTorch, segmentation_models_pytorch, UNet++"),
        ("Geospatial Libraries", "GDAL, Rasterio, Shapely, GeoPandas"),
        ("Spatial Database", "PostgreSQL 15 + PostGIS extension"),
        ("Backend Framework", "FastAPI (Python async microservice)"),
        ("Frontend Web GIS", "Leaflet.js / Mapbox GL JS, GeoJSON"),
        ("Deployment", "Docker microservices, AWS S3 / EC2")
    ]
    for head, body in col3_items:
        p = rtf3.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = f"•  {head}: "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_WHITE

        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    # Save output PPTX
    out_path = "GeoNex_SIH_2026_Presentation.pptx"
    prs.save(out_path)
    print(f"Presentation successfully created at: {os.path.abspath(out_path)}")

if __name__ == "__main__":
    create_deck()
